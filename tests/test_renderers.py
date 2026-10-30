"""Office and HTML renderers: robustness on model output and the safety of generated pages."""
import io
import re

from docx import Document
from pptx import Presentation

import config
import renderers


def test_pptx_tolerates_blank_cells_ragged_rows_and_no_bullets():
    data = renderers.build_pptx({"title": "Revue", "subtitle": "", "slides": [
        {"title": "Tableau", "bullets": ["a", ""], "table": {"columns": ["A", "B"], "rows": [["", None], ["seul"], ["1", "2", "3"]]}},
        {"title": "Graphique seul", "chart": {"type": "bar", "categories": ["L1", "L2"], "series": {"TRS": [80, "n/a"]}}},
        {"title": "Vide", "table": {"columns": ["X"], "rows": []}}]})
    prs = Presentation(io.BytesIO(data))
    texts = " ".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
    assert "[contenu]" not in texts
    assert prs.core_properties.keywords == "contenu-genere-par-ia"


def test_docx_8d_matches_headings_with_punctuation_and_marks_ai():
    data = renderers.build_docx({"template": "8D", "title": "8D", "sections": [
        {"heading": "D2: Description", "paragraphs": ["Quoi : défaut d'assemblage"]}, {"heading": "", "paragraphs": ["ignoré"]},
        {"heading": "d4. Causes", "bullets": ["Focale"], "table": {"columns": ["Cause", "Preuve"], "rows": [["Focale"]]}}]})
    doc = Document(io.BytesIO(data))
    body = "\n".join(p.text for p in doc.paragraphs)
    assert "Quoi : défaut d'assemblage" in body and "Focale" in body
    assert doc.core_properties.keywords == "contenu-genere-par-ia"


def test_docx_free_form_has_ai_footer():
    doc = Document(io.BytesIO(renderers.build_docx({"title": "Note", "sections": [{"heading": "Contexte", "paragraphs": [""]}]})))
    assert "généré par IA" in doc.sections[0].footer.paragraphs[0].text


def test_html_body_is_layout_only():
    body = ('<section class="em-panel" id="a"><div id="k" onclick="steal()" style="color:red;width:40%"></div>'
            '<script>alert(1)</script><iframe src="https://x"></iframe><a href="javascript:alert(1)">x</a>'
            '<img src=x onerror=alert(1)><form action="https://x"><input name=a></form></section>')
    html = renderers.build_html({"title": "t", "tabs": [{"id": "a", "label": "A"}], "body_html": body, "script": ""})
    main = html.split("<main>", 1)[1].split("</main>", 1)[0]
    assert not re.search(r"(?i)<script|<iframe|onclick|onerror|javascript:|<form|<input|href=|src=", main)
    assert 'style="width:40%"' in main


def test_html_data_cannot_close_the_script_tag():
    html = renderers.build_html({"title": "t", "data": {"x": "</script><img src=x onerror=alert(1)>"}, "script": ""})
    assert "</script><img" not in html and "\\u003c/script\\u003e" in html


def test_html_page_is_branded_and_marked_ai():
    html = renderers.build_html({"title": "Carte", "script": ""})
    assert config.COMPANY in html and 'data-ai-generated="true"' in html and config.COLORS["primary"].lower() in html.lower()
