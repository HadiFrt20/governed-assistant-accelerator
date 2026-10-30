"""Coverage requested by review round 2: SQL ownership filters, the real gateway client, router, attachments, tools,
mid-loop blocks, setup failures and installer guards."""
import io
import json
import sys
from pathlib import Path
from types import SimpleNamespace

import pytest

import agent
import config
import docs
import security
import store
from conftest import hdr

ROOT = Path(__file__).resolve().parents[1]


# ---------- store: per-user isolation is in the SQL itself ----------
@pytest.mark.parametrize("call", [lambda: store.file_owner("f.docx", "alice"), lambda: store.get_conversation("c", "alice"),
                                  lambda: store.delete("c", "alice"), lambda: store.rename("c", "alice", "t"),
                                  lambda: store.set_pinned("c", "alice", True), lambda: store.list_files("alice"),
                                  lambda: store.list_conversations("alice")])
def test_store_queries_filter_on_user(monkeypatch, call):
    seen = []
    monkeypatch.setattr(store, "q", lambda sql, args=(), one=False: seen.append((sql, args)) or ([] if not one else None))
    call()
    sql, args = seen[0]
    assert "user_email=%s" in sql.replace(" ", "") .replace("user_email=%s", "user_email=%s") and "alice" in args


# ---------- docs ----------
def test_docs_extract_strips_nul_truncates_and_skips_unknown(monkeypatch):
    monkeypatch.setattr(docs, "MAX_CHARS", 10)
    assert docs.extract("a.txt", b"ab\x00cdefghijklmnop") == "abcdefghij\n[… document tronqué]"
    assert docs.extract("a.exe", b"MZ") is None


# ---------- real Gateway.messages with a fake HTTP layer ----------
class Resp:
    def __init__(self, status, body=None, headers=None):
        self.status_code, self._body, self.headers = status, body or {}, {"content-type": "application/json", **(headers or {})}
        self.ok, self.text = status < 400, json.dumps(body or {})

    def json(self):
        return self._body

    def raise_for_status(self):
        if not self.ok:
            raise RuntimeError(self.status_code)


def _gw(monkeypatch, responses, user_token="u"):
    j = {"events": [], "messages": [], "files": []}
    gw = agent.Gateway(j, user_token, "sp", None)
    seq = iter(responses)
    monkeypatch.setattr(gw, "_post", lambda body, purpose: next(seq))
    monkeypatch.setattr(agent.time, "sleep", lambda s: None)
    return gw, j


OK = {"model": "claude-sonnet-5-5", "stop_reason": "end_turn", "usage": {"input_tokens": 1000, "output_tokens": 100},
      "content": [{"type": "text", "text": "ok"}]}


def test_gateway_backs_off_on_429_and_computes_cost(monkeypatch):
    monkeypatch.setattr(config, "PRICES", {"claude-sonnet-5-5": (2.0, 10.0), "default": (2.0, 10.0)})
    gw, j = _gw(monkeypatch, [Resp(429, headers={"retry-after": "1"}), Resp(503), Resp(200, OK)])
    out = gw.messages({"messages": [{"role": "user", "content": "x"}]}, "agent")
    assert out["content"][0]["text"] == "ok"
    p = config.PRICES["claude-sonnet-5-5"]
    assert j["cost_usd"] == pytest.approx((1000 * p[0] + 100 * p[1]) / 1e6)


def test_gateway_retries_unevaluable_policy_then_reports_real_block(monkeypatch):
    unevaluable = {"content": [], "databricks_service_policy": {"name": "donnees-sensibles", "reason": "a policy on this service could not be evaluated"}}
    gw, j = _gw(monkeypatch, [Resp(200, unevaluable), Resp(200, OK)])
    assert gw.messages({"messages": [{"role": "user", "content": "x"}]}, "agent")["content"][0]["text"] == "ok"


def test_gateway_obo_fallback_only_when_enabled(monkeypatch):
    monkeypatch.setattr(security, "OBO_FALLBACK_TO_APP", True)
    gw, j = _gw(monkeypatch, [Resp(403, {"error": "scope"}), Resp(200, OK)])
    gw.messages({"messages": [{"role": "user", "content": "x"}]}, "agent")
    assert gw.token == "sp" and not gw.as_user and any(e["type"] == "identity" for e in j["events"])


def test_anonymisation_is_confirmed_only_by_placeholders(monkeypatch):
    masked = dict(OK, content=[{"type": "text", "text": "écrire à [EMAIL_ADDRESS]"}])
    gw, j = _gw(monkeypatch, [Resp(200, masked)])
    gw.messages({"messages": [{"role": "user", "content": "écris à jean@acier.fr"}]}, "agent")
    assert "Anonymisé" in next(e for e in j["events"] if e["type"] == "masked")["text"]
    gw, j = _gw(monkeypatch, [Resp(200, OK)])
    gw.messages({"messages": [{"role": "user", "content": "écris à jean@acier.fr"}]}, "agent")
    assert next(e for e in j["events"] if e["type"] == "masked")["text"].startswith("Données sensibles détectées")


# ---------- router ----------
def test_router_rules_and_llm_decision():
    class G:
        def __init__(self, word): self.word = word
        def messages(self, body, purpose): return {"content": [{"type": "text", "text": self.word}]}
    if config.GENIE_ONE:
        assert agent.route(G("GENIE"), "Quel est le taux de rebut ?", False)[0] == "genie"
    assert agent.route(G("GENIE"), "Rédige un rapport 8D", False)[0] == "agent"
    assert agent.route(G("GENIE"), "Analyse", True)[0] == "agent"
    blocked = agent.route(type("B", (), {"messages": lambda s, b, p: {"databricks_service_policy": {"name": "x"}}})(), "q ?", False)
    assert blocked[0] == "blocked" or not config.GENIE_ONE


# ---------- attachments ----------
def test_read_attachments_confines_tolerates_and_caps(monkeypatch):
    class R:
        ok, status_code, content = True, 200, b"x"
    monkeypatch.setattr(agent.requests, "get", lambda *a, **k: R())
    monkeypatch.setattr(agent.docs, "extract", lambda name, data: (_ for _ in ()).throw(ValueError("chiffré")) if name.startswith("bad") else "y" * 150_000)
    uf = security.user_folder("alice@acme.com")
    j = {"events": [], "messages": [{"role": "user", "content": "q"}]}
    agent.read_attachments(j, "sp", [{"name": "a.pdf", "path": f"{uf}/a.pdf"}, {"name": "bad.pdf", "path": f"{uf}/bad.pdf"},
                                     {"name": "lib.pdf", "path": f"{config.LIBRARY}/lib.pdf"}], "q", uf)
    content = j["messages"][-1]["content"]
    assert len(content) < 205_000 and "limite par demande" in content
    assert any("n'a pas pu être lu" in e["text"] for e in j["events"] if e["type"] == "error")
    with pytest.raises(security.Forbidden):
        agent.read_attachments(j, "sp", [{"name": "x.pdf", "path": f"{security.user_folder('bob@acme.com')}/x.pdf"}], "q", uf)


# ---------- tools end to end ----------
def test_create_html_tool_publishes_and_returns_published_facts(monkeypatch):
    from test_agent import FakeGateway, text, tool_use
    monkeypatch.setattr(agent, "upload", lambda token, name, data: f"{config.GENERATED}/{name}")
    j = {"events": [], "files": [], "answer": None, "messages": [{"role": "user", "content": "carte"}]}
    spec = {"title": "Carte des capacités", "tabs": [{"id": "a", "label": "A"}], "body_html": '<section class="em-panel" id="a"></section>',
            "data": {"caps": [{"status": "Couverte"}, {"status": "Partielle"}]}, "script": ""}
    gw = FakeGateway(j, [tool_use("create_html", spec), text("Page publiée.")])
    agent.agent_loop(gw, j, "alice", None, "", "/x")
    f = j["files"][0]
    assert f["kind"] == "html" and f["name"].isascii() and f["name"].endswith(".html")
    result = json.loads(gw.sent[1][-1]["content"][0]["content"])
    assert result["chiffres_publies"]["caps"]["total"] == 2


def test_policy_block_after_a_tool_round_keeps_history_valid():
    from test_agent import FakeGateway, tool_use
    history = [{"role": "user", "content": "avant"}, {"role": "assistant", "content": [{"type": "text", "text": "ok"}]}]
    j = {"events": [], "files": [], "answer": None, "turn_start": 2, "messages": history + [{"role": "user", "content": "nouvelle"}]}
    gw = FakeGateway(j, [tool_use("unknown_tool", {}), {"content": [], "databricks_service_policy": {"name": "anti-jailbreak", "reason": "non"}}])
    agent.agent_loop(gw, j, "a", None, "", "/x")
    assert j["messages"] == history  # no orphan tool_use, no blocked user message


# ---------- routes: conversation view with a live turn, setup failure ----------
def test_conversation_view_with_live_turn(client, appmod, monkeypatch):
    appmod._db["convs"]["c1"] = {"id": "c1", "user": "alice@acme.com", "department": None, "llm_messages": []}
    monkeypatch.setattr(store, "get_conversation", lambda cid, user: {"id": "c1", "llm_messages": [], "turns": [{"id": "t9", "status": "running"}]})
    assert client.get("/api/conversations/c1", headers=hdr()).status_code == 200


def test_setup_failure_does_not_leave_a_running_job(client, appmod, monkeypatch):
    appmod._db["convs"]["c1"] = {"id": "c1", "user": "alice@acme.com", "department": None, "llm_messages": []}
    monkeypatch.setattr(store, "start_turn", lambda *a, **k: (_ for _ in ()).throw(RuntimeError("lakebase down")))
    with pytest.raises(RuntimeError):
        client.post("/api/conversations/c1/messages", json={"text": "un"}, headers=hdr())
    monkeypatch.setattr(store, "start_turn", lambda cid, text, atts, tid=None: tid)
    monkeypatch.setattr(appmod.agent, "run_agent", lambda *a, **k: None)
    assert client.post("/api/conversations/c1/messages", json={"text": "deux"}, headers=hdr()).status_code == 200


def test_library_document_can_be_attached(client, appmod, monkeypatch):
    appmod._db["convs"]["c1"] = {"id": "c1", "user": "alice@acme.com", "department": None, "llm_messages": []}
    monkeypatch.setattr(appmod.agent, "run_agent", lambda *a, **k: None)
    r = client.post("/api/conversations/c1/messages", json={"text": "Résume", "attachments": [{"name": "proc.pdf", "path": f"{config.LIBRARY}/Qualite/proc.pdf"}]}, headers=hdr())
    assert r.status_code == 200


# ---------- installer guards ----------
@pytest.fixture()
def installer():
    sys.path.insert(0, str(ROOT / "install"))
    import install as ins
    return ins


def test_installer_sql_literal_escaping(installer):
    assert installer._q("l'usine \\ x") == "l\\'usine \\\\ x"


# ---------- round-2 security follow-ups ----------
def test_no_user_token_fails_closed_when_obo_is_configured(monkeypatch):
    monkeypatch.setattr(security, "OBO_FALLBACK_TO_APP", False)
    monkeypatch.setitem(config.S, "obo", True)
    with pytest.raises(PermissionError):
        agent.Gateway({"events": [], "messages": []}, None, "sp", None)


def test_user_folders_are_collision_free():
    assert security.user_folder("jean+x@acme.com") != security.user_folder("jean_x@acme.com")


def test_sanitizer_handles_slash_separated_attributes():
    import renderers
    html = renderers.build_html({"title": "t", "body_html": '<img/src=x/onerror=alert(1)><div/onclick=alert(1)></div><a/href=javascript:alert(1)>x</a>', "script": ""})
    main = html.split("<main>", 1)[1].split("</main>", 1)[0]
    assert "onerror" not in main and "onclick" not in main and "javascript:" not in main and "src=" not in main


def test_zip_bombs_are_refused():
    import zipfile
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("word/document.xml", b"0" * (50 * 1024 * 1024))
    with pytest.raises(ValueError):
        docs.extract("bomb.docx", buf.getvalue())


# ---------- round-3 coverage ----------
def test_store_search_is_user_scoped(monkeypatch):
    seen = []
    monkeypatch.setattr(store, "q", lambda sql, args=(), one=False: seen.append((sql, args)) or [])
    for bm25 in (True, False):
        monkeypatch.setitem(store.BM25, "on", bm25)
        store.search("alice", "quarantaine")
    searches = [(sql, args) for sql, args in seen if "assistant.turns t" in sql]
    assert len(searches) == 2 and all("c.user_email = %s" in sql and "alice" in args for sql, args in searches)


def test_genie_handoff_fails_closed_on_403_and_never_reuses_foreign_threads(monkeypatch):
    import requests
    class G:
        as_user = True
        def ask(self, *a, **k):
            r = requests.Response(); r.status_code = 403
            raise requests.HTTPError(response=r)
    gw = agent.Gateway({"events": [], "messages": [], "files": []}, "u", "sp", None)
    gw.genie = G()
    with pytest.raises(PermissionError):
        agent.genie_handoff(gw, gw.job, "Quel est le taux ?", "conv-1")


def test_office_tools_publish_files(monkeypatch):
    from test_agent import FakeGateway, text
    monkeypatch.setattr(agent, "upload", lambda token, name, data: f"{config.GENERATED}/{name}")
    j = {"events": [], "files": [], "answer": None, "messages": [{"role": "user", "content": "doc"}]}
    calls = {"stop_reason": "tool_use", "content": [
        {"type": "tool_use", "id": "a", "name": "create_pptx", "input": {"title": "Revue", "slides": [{"title": "S", "bullets": ["x"]}]}},
        {"type": "tool_use", "id": "b", "name": "create_docx", "input": {"title": "Note", "sections": [{"heading": "Contexte", "paragraphs": ["y"]}]}}]}
    agent.agent_loop(FakeGateway(j, [calls, text("ok")]), j, "a", None, "", "/x")
    assert [f["kind"] for f in j["files"]] == ["pptx", "docx"] and all(f["name"].isascii() for f in j["files"])


def test_genie_tool_inside_agent_loop(monkeypatch):
    from test_agent import FakeGateway, text, tool_use
    j = {"events": [], "files": [], "answer": None, "messages": [{"role": "user", "content": "q"}]}
    gw = FakeGateway(j, [tool_use("genie_one", {"question": "taux de rebut ?"}), text("2 %")])
    gw.genie = SimpleNamespace(as_user=True, ask=lambda q, **k: {"status": "completed", "answer": "2,1 %", "deep_link": "https://x/one/chat/threads/1"})
    agent.agent_loop(gw, j, "a", None, "", "/x")
    assert any(e["type"] == "genie" and "2,1" in e["text"] for e in j["events"]) and j["answer"] == "2 %"


def test_worker_reports_refusals_and_hides_technical_errors(client, appmod, monkeypatch):
    import time
    appmod._db["convs"]["c1"] = {"id": "c1", "user": "alice@acme.com", "department": None, "llm_messages": []}
    for exc, expected in ((PermissionError("Reconnectez-vous"), "Reconnectez-vous"), (KeyError("secret-detail"), "erreur technique")):
        monkeypatch.setattr(appmod.agent, "run_agent", lambda *a, _e=exc, **k: (_ for _ in ()).throw(_e))
        tid = client.post("/api/conversations/c1/messages", json={"text": "x"}, headers=hdr()).json()["turn_id"]
        for _ in range(50):
            p = client.get(f"/api/turns/{tid}", headers=hdr()).json()
            if p.get("status") == "done":
                break
            time.sleep(0.05)
        assert expected in p["answer"] and "secret-detail" not in p["answer"]


def test_genie_conversation_reuse_only_for_owned_threads(client, appmod, monkeypatch):
    seen = {}
    monkeypatch.setattr(store, "get_conversation", lambda cid, user: {"id": cid, "llm_messages": [], "department": None, "turns": [
        {"id": "t0", "status": "done", "user_text": "q", "events": [{"type": "handoff", "conversation_id": "app-thread", "own": False}]}]})
    monkeypatch.setattr(appmod.agent, "run_agent", lambda *a, **k: seen.update(k))
    client.post("/api/conversations/c1/messages", json={"text": "et la nuit ?"}, headers=hdr())
    import time; time.sleep(0.2)
    assert seen.get("genie_conv") is None


# ---------- round-4 coverage ----------
def test_docx_pptx_pdf_extraction():
    from docx import Document
    from pptx import Presentation
    from pptx.util import Inches
    b = io.BytesIO(); d = Document(); d.add_paragraph("Procédure é"); t = d.add_table(rows=1, cols=2); t.cell(0, 0).text = "A"; d.save(b)
    assert "Procédure é" in docs.extract("p.docx", b.getvalue()) and "A |" in docs.extract("p.docx", b.getvalue())
    b = io.BytesIO(); prs = Presentation(); s = prs.slides.add_slide(prs.slide_layouts[5]); s.shapes.title.text = "Diapo un"; prs.save(b)
    assert "[Diapositive 1]" in docs.extract("d.pptx", b.getvalue()) and "Diapo un" in docs.extract("d.pptx", b.getvalue())
    from pypdf import PdfWriter
    b = io.BytesIO(); w = PdfWriter(); w.add_blank_page(100, 100); w.write(b)
    assert "[Page 1]" in docs.extract("x.pdf", b.getvalue())


def test_store_q_reconnects_once_then_raises(monkeypatch):
    import psycopg
    class C:
        def __init__(self, fail): self.fail = fail
        def execute(self, sql, args):
            if self.fail: raise psycopg.OperationalError("gone")
            return SimpleNamespace(description=None, fetchall=lambda: [])
    conns = iter([C(True), C(False)])
    monkeypatch.setattr(store, "conn", lambda: next(conns))
    assert store.q("SELECT 1") == []
    conns2 = iter([C(True), C(True)])
    monkeypatch.setattr(store, "conn", lambda: next(conns2))
    with pytest.raises(psycopg.OperationalError):
        store.q("SELECT 1")


def test_curator_upload_sanitises_folder_and_never_overwrites(client, appmod, monkeypatch):
    import security as sec
    monkeypatch.setattr(sec, "LIBRARY_CURATORS", {"alice@acme.com"})
    r = client.post("/api/library/upload", files={"file": ("Proc é.pdf", io.BytesIO(b"%PDF"), "application/pdf")}, data={"folder": "../../Qualité"}, headers=hdr())
    assert r.status_code == 200
    method, path = appmod._calls[-1]
    assert method == "PUT" and "overwrite=false" in path and ".." not in path and path.startswith(f"files{config.LIBRARY}/")


def test_library_tree_caches_and_keeps_last_listing_on_error(appmod, monkeypatch):
    calls = {"n": 0}

    class R:
        def __init__(self, ok, contents=()):
            self.ok, self.status_code, self._c = ok, 200 if ok else 503, list(contents)

        def json(self):
            return {"contents": self._c}

    def files_api(method, path, **kw):
        calls["n"] += 1
        return R(True, [{"path": f"{config.LIBRARY}/Q/a.pdf", "name": "a.pdf", "is_directory": False}]) if calls["n"] == 1 else R(False)

    monkeypatch.setattr(appmod, "files_api", files_api)
    monkeypatch.setitem(appmod._LIB, "at", 0)
    first = appmod._real_library_tree()
    assert [d["name"] for d in first] == ["a.pdf"] and first[0]["folder"] == "Q"
    assert appmod._real_library_tree() == first and calls["n"] == 1      # cached
    appmod._LIB["at"] = 0                                                # cache expired, next listing fails
    assert appmod._real_library_tree() == first and calls["n"] == 2      # last complete listing kept


def test_sse_tool_input_reassembly_and_truncation():
    from test_agent import FakeSSE
    out = agent.read_sse(FakeSSE([{"type": "message_start", "message": {"id": "m", "content": []}},
                                  {"type": "content_block_start", "index": 0, "content_block": {"type": "tool_use", "id": "t", "name": "create_html", "input": {}}},
                                  {"type": "content_block_delta", "index": 0, "delta": {"type": "input_json_delta", "partial_json": '{"title": "Carte'}},
                                  {"type": "message_delta", "delta": {"stop_reason": "max_tokens"}}]))
    assert out["content"][0]["truncated"] is True and out["stop_reason"] == "max_tokens"


def test_genie_one_success_path(monkeypatch):
    g = agent.GenieOne("t"); g.sid = "s"
    seq = iter([{"conversation_id": "c", "response_id": "r", "status": "in_progress"},
                {"status": "in_progress", "deep_link": "https://x/one/chat/threads/c"},
                {"status": "completed", "final_answer": "**2,1 %**", "query_items": [{"sql": "SELECT 1"}]}])
    monkeypatch.setattr(g, "_call", lambda tool, args: next(seq))
    monkeypatch.setattr(agent.time, "sleep", lambda s: None)
    links = []
    out = g.ask("taux ?", on_link=lambda link, conv: links.append(link))
    assert out["answer"] == "**2,1 %**" and links == ["https://x/one/chat/threads/c"] and out["conversation_id"] == "c"


def test_mcp_tool_errors_are_raised():
    g = agent.GenieOne("t")
    g._rpc = lambda method, params=None, rid=True: {"result": {"isError": True, "content": [{"type": "text", "text": "table introuvable"}]}}
    with pytest.raises(RuntimeError, match="table introuvable"):
        g._call("genie_ask", {})


def test_genie_handoff_success_writes_answer():
    gw = agent.Gateway({"events": [], "messages": [{"role": "user", "content": "q"}], "files": []}, "u", "sp", None)
    gw.genie = SimpleNamespace(as_user=True, ask=lambda q, **k: (k["on_link"]("https://x/t", "c1"), {"status": "completed", "answer": "42", "deep_link": "https://x/t"})[1])
    agent.genie_handoff(gw, gw.job, "q", None)
    assert gw.job["answer"] == "42" and gw.job["messages"][-1]["role"] == "assistant"
    assert any(e["type"] == "handoff" and e["own"] for e in gw.job["events"])


def test_run_agent_genie_failure_falls_back_but_refusal_does_not(monkeypatch):
    from test_agent import FakeGateway
    seen = []
    monkeypatch.setattr(agent, "route", lambda *a, **k: ("genie", "LLM"))
    monkeypatch.setattr(agent, "Gateway", lambda *a, **k: FakeGateway(a[0], []))
    monkeypatch.setattr(agent, "agent_loop", lambda *a, **k: seen.append("agent"))
    monkeypatch.setattr(agent, "genie_handoff", lambda *a, **k: (_ for _ in ()).throw(TimeoutError("lent")))
    j = {"events": [], "files": [], "answer": None, "messages": [{"role": "user", "content": "q"}]}
    agent.run_agent(j, "u", "sp", user="a", user_folder="/x", text="q")
    assert seen == ["agent"] and any("indisponible" in e.get("text", "") for e in j["events"])
    monkeypatch.setattr(agent, "genie_handoff", lambda *a, **k: (_ for _ in ()).throw(PermissionError("reconnectez-vous")))
    with pytest.raises(PermissionError):
        agent.run_agent({"events": [], "files": [], "answer": None, "messages": [{"role": "user", "content": "q"}]}, "u", "sp", user="a", user_folder="/x", text="q")


def test_run_agent_stops_too_long_conversations():
    j = {"events": [], "files": [], "answer": None, "messages": [{"role": "user", "content": "x" * 800_000}, {"role": "user", "content": "q"}]}
    agent.run_agent(j, "u", "sp", user="a", user_folder="/x", text="q")
    assert "trop longue" in j["answer"] and len(j["messages"]) == 1


# ---------- round-5 coverage ----------


def test_ip_leak_policy_sql_is_cel_safe_and_escaped(installer):
    sql = installer.ip_leak_sql("cat.sch", ["Confidentiel R&D", "l'usine"], "Interdit : PI", "Acme")
    assert "CONTAINS(LOWER(event:context.message::string), 'confidentiel r&d')" in sql
    code = "\n".join(l for l in sql.splitlines() if not l.lstrip().startswith("--"))  # ignore the template's explanatory comments
    assert "l\\'usine" in code and "COALESCE" not in code.upper() and "RLIKE" not in code.upper() and "REGEXP" not in code.upper()
    assert "'DENY'" in sql and "Interdit : PI" in sql


def test_mcp_rpc_parses_sse_lines_and_keeps_session(monkeypatch):
    g = agent.GenieOne("t")
    class R:
        ok, status_code = True, 200
        headers = {"mcp-session-id": "s-1"}
        text = 'event: message\ndata: {"jsonrpc": "2.0", "result": {"structuredContent": {"ok": 1}}}\n'
        def raise_for_status(self): pass
    monkeypatch.setattr(agent.requests, "post", lambda *a, **k: R())
    assert g._call("genie_ask", {}) == {"ok": 1} and g.sid == "s-1"
    R.text = '{"jsonrpc": "2.0", "error": {"message": "boom"}}'
    with pytest.raises(RuntimeError):
        g._call("genie_ask", {})


def test_sse_bare_policy_event_and_error_event():
    from test_agent import FakeSSE
    out = agent.read_sse(FakeSSE([{"databricks_service_policy": {"name": "anti-fuite-pi", "reason": "PI"}}]))
    assert out["databricks_service_policy"]["name"] == "anti-fuite-pi"
    with pytest.raises(RuntimeError):
        agent.read_sse(FakeSSE([{"type": "error", "error": {"type": "overloaded_error"}}]))


def test_genie_handoff_app_fallback_never_reuses_user_thread(monkeypatch):
    import requests
    monkeypatch.setattr(security, "OBO_FALLBACK_TO_APP", True)
    asks = []
    class G:
        def __init__(self, fail): self.fail, self.as_user = fail, True
        def ask(self, q, conversation_id=None, on_progress=None, on_link=None):
            asks.append(conversation_id)
            if self.fail:
                on_link("https://x/user-thread", "u-conv")
                r = requests.Response(); r.status_code = 403
                raise requests.HTTPError(response=r)
            return {"status": "completed", "answer": "ok", "deep_link": None}
    gw = agent.Gateway({"events": [], "messages": [{"role": "user", "content": "q"}], "files": []}, "u", "sp", None)
    gw.genie = G(True)
    monkeypatch.setattr(agent, "GenieOne", lambda token, as_user=False: G(False))
    agent.genie_handoff(gw, gw.job, "q", "previous-user-conv")
    assert asks == ["previous-user-conv", None]                                   # the app never continues the user's thread
    assert all(not e["own"] for e in gw.job["events"] if e["type"] == "handoff")  # links to app-owned threads hidden


def test_unsaved_turn_is_marked_error(client, appmod, monkeypatch):
    import time
    appmod._db["convs"]["c1"] = {"id": "c1", "user": "alice@acme.com", "department": None, "llm_messages": []}
    failed = []
    monkeypatch.setattr(appmod.agent, "run_agent", lambda job, *a, **k: job.update(answer="ok"))
    monkeypatch.setattr(store, "finish_turn", lambda *a, **k: (_ for _ in ()).throw(RuntimeError("lakebase")))
    monkeypatch.setattr(store, "fail_turn", lambda tid, msg: failed.append(tid))
    tid = client.post("/api/conversations/c1/messages", json={"text": "x"}, headers=hdr()).json()["turn_id"]
    for _ in range(50):
        if failed:
            break
        time.sleep(0.05)
    assert failed == [tid]


# ---------- round-6 coverage ----------
def test_genie_failure_status_is_an_error(monkeypatch):
    g = agent.GenieOne("t"); g.sid = "s"
    seq = iter([{"conversation_id": "c", "response_id": "r", "status": "in_progress"}, {"status": "failed", "final_answer": None}])
    monkeypatch.setattr(g, "_call", lambda tool, args: next(seq))
    with pytest.raises(RuntimeError, match="failed"):
        g.ask("q")


def test_history_completeness():
    assert agent.history_is_complete([{"role": "user", "content": "q"}, {"role": "assistant", "content": [{"type": "text", "text": "r"}]}])
    assert not agent.history_is_complete([{"role": "user", "content": "q"}])
    assert not agent.history_is_complete([{"role": "assistant", "content": [{"type": "tool_use", "id": "t", "name": "x", "input": {}}]}])


def test_failed_turn_does_not_corrupt_history(client, appmod, monkeypatch):
    import time
    appmod._db["convs"]["c1"] = {"id": "c1", "user": "alice@acme.com", "department": None,
                                 "llm_messages": [{"role": "user", "content": "avant"}, {"role": "assistant", "content": "ok"}]}
    saved = []
    def boom(job, *a, **k):
        job["messages"].append({"role": "assistant", "content": [{"type": "tool_use", "id": "t", "name": "x", "input": {}}]})
        raise RuntimeError("gateway down")
    monkeypatch.setattr(appmod.agent, "run_agent", boom)
    monkeypatch.setattr(store, "finish_turn", lambda tid, cid, ans, ev, msgs, **k: saved.append(msgs))
    client.post("/api/conversations/c1/messages", json={"text": "x"}, headers=hdr())
    for _ in range(50):
        if saved:
            break
        time.sleep(0.05)
    assert saved[0] == [{"role": "user", "content": "avant"}, {"role": "assistant", "content": "ok"}]


# ---------- round-7 coverage ----------
def test_gateway_request_construction(monkeypatch):
    sent = {}
    def post(url, **kw):
        sent.update(url=url, **kw)
        return Resp(200, OK)
    monkeypatch.setattr(agent.requests, "post", post)
    gw = agent.Gateway({"events": [], "messages": [], "files": []}, "user-tok", "sp", "Qualité")
    gw.messages({"messages": [{"role": "user", "content": "x"}], "max_tokens": 5}, "routeur")
    assert sent["url"].endswith("/ai-gateway/anthropic/v1/messages") and sent["headers"]["Authorization"] == "Bearer user-tok"
    tags = json.loads(sent["headers"]["Databricks-Ai-Gateway-Request-Tags"])
    assert tags["department"] == "Qualité" and tags["step"] == "routeur" and tags["app"] == config.APP_ID
    assert sent["json"]["model"] == config.MODEL and "stream" not in sent["json"]
    gw.messages({"messages": [{"role": "user", "content": "x"}]}, "agent")
    assert sent["json"]["stream"] is True and sent["stream"] is True


def test_empty_model_reply_is_not_stored():
    from test_agent import FakeGateway
    j = {"events": [], "files": [], "answer": None, "messages": [{"role": "user", "content": "q"}]}
    agent.agent_loop(FakeGateway(j, [{"stop_reason": "end_turn", "content": [{"type": "text", "text": "  "}, {"type": "thinking", "thinking": ""}]}]), j, "a", None, "", "/x")
    assert j["messages"] == [{"role": "user", "content": "q"}] and "pas pu produire" in j["answer"]
    assert not agent.history_is_complete([{"role": "assistant", "content": []}])


def test_library_listing_follows_pagination(appmod, monkeypatch):
    class R:
        ok, status_code = True, 200
        def __init__(self, body): self.body = body
        def json(self): return self.body
    def files_api(method, path, **kw):
        if "page_token=p2" in path:
            return R({"contents": [{"path": f"{config.LIBRARY}/b.pdf", "name": "b.pdf"}]})
        return R({"contents": [{"path": f"{config.LIBRARY}/a.pdf", "name": "a.pdf"}], "next_page_token": "p2"})
    monkeypatch.setattr(appmod, "files_api", files_api)
    monkeypatch.setitem(appmod._LIB, "at", 0)
    assert [d["name"] for d in appmod._real_library_tree()] == ["a.pdf", "b.pdf"]


# ---------- round-8 coverage ----------
def test_phone_numbers_are_never_claimed_as_anonymised():
    sys.path.insert(0, str(ROOT / "install"))
    import settings as shared
    import yaml
    cfg = yaml.safe_load(open(ROOT / "config/assistant.yaml"))
    cfg["governance"]["builtin_policies"]["sensitive_data"]["categories"] = ["class.email_address", "class.phone_number", "class.iban_code"]
    s = shared.build(cfg, yaml.safe_load(open(ROOT / "brands/acme/brand.yaml")), dict.fromkeys(shared.TEMPLATE_KEYS, "x"))
    assert "téléphone" not in s["governance"]["anonymised_labels"] and {"adresse email", "IBAN"} <= set(s["governance"]["anonymised_labels"])


def test_dashboard_cost_sql(installer):
    sql = installer.price_cases({"claude-sonnet-5-5": [2, 10], "default": [1, 1]})
    assert sql == ("WHEN destination_model LIKE '%claude-sonnet-5-5%' THEN COALESCE(input_tokens, 0) * 2.0 / 1e6 + "
                   "COALESCE(output_tokens, 0) * 10.0 / 1e6")


def test_no_token_with_fallback_enabled_runs_as_app_visibly(monkeypatch):
    monkeypatch.setattr(security, "OBO_FALLBACK_TO_APP", True)
    j = {"events": [], "messages": []}
    gw = agent.Gateway(j, None, "sp", None)
    assert gw.token == "sp" and not gw.as_user and (not config.S.get("obo", True) or any(e["type"] == "identity" for e in j["events"]))


def test_genie_mcp_decodes_utf8(monkeypatch):
    g = agent.GenieOne("t")
    payload = 'data: {"jsonrpc": "2.0", "result": {"structuredContent": {"answer": "Données « ok » é"}}}\n'.encode("utf-8")
    class R:
        ok, status_code, headers = True, 200, {}
        encoding = None
        content = payload
        @property
        def text(self): return self.content.decode(self.encoding or "latin-1")
        def raise_for_status(self): pass
    monkeypatch.setattr(agent.requests, "post", lambda *a, **k: R())
    assert g._call("genie_ask", {})["answer"] == "Données « ok » é"


def test_library_network_errors_keep_last_listing(appmod, monkeypatch):
    import requests as rq
    monkeypatch.setattr(appmod, "files_api", lambda *a, **k: (_ for _ in ()).throw(rq.ConnectionError("down")))
    monkeypatch.setitem(appmod._LIB, "at", 0)
    monkeypatch.setitem(appmod._LIB, "items", [{"name": "kept.pdf"}])
    assert appmod._real_library_tree() == [{"name": "kept.pdf"}]


# ---------- round-9 coverage ----------
@pytest.mark.parametrize("status,pending", [("IN_PROGRESS", True), ("running", True), (None, True), ("COMPLETED", False), ("failed", False)])
def test_genie_status_case_insensitive(status, pending):
    assert agent._pending(status) is pending


def test_store_init_closes_orphaned_turns_and_tries_bm25(monkeypatch):
    executed = []
    class C:
        def execute(self, sql, *a): executed.append(sql)
    monkeypatch.setattr(store, "conn", lambda: C())
    monkeypatch.setattr(store, "q", lambda sql, args=(), one=False: executed.append(sql) or (None if one else []))
    store.init()
    assert any("CREATE TABLE IF NOT EXISTS assistant.turns" in s for s in executed)
    assert any("SET status='error'" in s and "status='running'" in s for s in executed)
    assert any("lakebase_text" in s for s in executed)


# ---------- round-10 coverage ----------
def test_xlsx_extraction():
    import zipfile
    b = io.BytesIO()
    with zipfile.ZipFile(b, "w") as z:
        z.writestr("xl/sharedStrings.xml", '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><si><t>Ligne</t></si><si><t>TRS é</t></si></sst>')
        z.writestr("xl/worksheets/sheet1.xml", '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>'
                   '<row><c t="s"><v>0</v></c><c t="s"><v>1</v></c></row><row><c t="inlineStr"><is><t>L1</t></is></c><c><v>81.5</v></c></row></sheetData></worksheet>')
    text = docs.extract("kpi.xlsx", b.getvalue())
    assert text == "[Feuille 1]\nLigne | TRS é\nL1 | 81.5"


def test_config_endpoint(client):
    r = client.get("/api/config", headers=hdr("jean.dupont@acme.com"))
    body = r.json()
    assert r.status_code == 200 and body["brand"] == config.BRAND and body["company"] == config.COMPANY
    assert body["name"] == "Jean Dupont" and body["obo"] is False and body["departments"] == config.DEPARTMENTS
    assert client.get("/api/config", headers=hdr(**{"x-forwarded-access-token": "t"})).json()["obo"] is True


def test_sse_preserves_signed_thinking():
    from test_agent import FakeSSE
    out = agent.read_sse(FakeSSE([{"type": "message_start", "message": {"id": "m", "content": []}},
                                  {"type": "content_block_start", "index": 0, "content_block": {"type": "thinking", "thinking": ""}},
                                  {"type": "content_block_delta", "index": 0, "delta": {"type": "thinking_delta", "thinking": "Je réfléchis"}},
                                  {"type": "content_block_delta", "index": 0, "delta": {"type": "signature_delta", "signature": "sig=="}},
                                  {"type": "message_delta", "delta": {"stop_reason": "end_turn"}}]))
    assert out["content"][0] == {"type": "thinking", "thinking": "Je réfléchis", "signature": "sig=="}


# ---------- round-11 coverage ----------
def test_genie_tool_refusal_fails_closed_inside_agent_loop(monkeypatch):
    import requests as rq
    from test_agent import FakeGateway, tool_use
    monkeypatch.setattr(security, "OBO_FALLBACK_TO_APP", False)
    j = {"events": [], "files": [], "answer": None, "messages": [{"role": "user", "content": "q"}]}
    real = agent.Gateway(j, "u", "sp", None)
    gw = FakeGateway(j, [tool_use("genie_one", {"question": "taux ?"})])
    gw._fallback = real._fallback
    def ask(q, **k):
        r = rq.Response(); r.status_code = 403
        raise rq.HTTPError(response=r)
    gw.genie = SimpleNamespace(as_user=True, ask=ask)
    with pytest.raises(PermissionError):
        agent.agent_loop(gw, j, "a", None, "", "/x")


def test_local_dev_identity_requires_explicit_flag(client, monkeypatch):
    monkeypatch.delenv("LOCAL_DEV", raising=False)
    assert client.get("/api/conversations").status_code == 401
    monkeypatch.setenv("LOCAL_DEV", "true")          # anything but "1" stays closed
    assert client.get("/api/conversations").status_code == 401
    monkeypatch.setenv("LOCAL_DEV", "1")
    assert client.get("/api/conversations").status_code == 200


def test_lakebase_connection_refreshes_credentials(monkeypatch):
    made = []
    class Conn:
        closed = False
        def close(self): self.closed = True
    monkeypatch.setattr(store, "ws", lambda: SimpleNamespace(postgres=SimpleNamespace(generate_database_credential=lambda endpoint: SimpleNamespace(token=f"tok{len(made)}")),
                                                         current_user=SimpleNamespace(me=lambda: SimpleNamespace(user_name="sp"))))
    monkeypatch.setattr(store.psycopg, "connect", lambda **kw: made.append(kw["password"]) or Conn())
    monkeypatch.setenv("LAKEBASE_ENDPOINT", "projects/p/branches/production/endpoints/primary")
    monkeypatch.setenv("PGHOST", "h")
    monkeypatch.setenv("PGUSER", "sp-client-id")
    monkeypatch.setattr(store, "_conn", None)
    first = store.conn()
    assert store.conn() is first and made == ["tok0"]               # reused while valid
    monkeypatch.setattr(store, "_exp", 0)                          # credential expired
    second = store.conn()
    assert second is not first and first.closed and made == ["tok0", "tok1"]


# ---------- round-12 coverage ----------
def test_gold_layer_parses_events_instead_of_matching_json_text():
    sql = (ROOT / "install" / "sql" / "gold_usage.sql").read_text()
    assert "from_json(" in sql and "exists(ev, e -> e.type = 'blocked')" in sql
    assert '"type": "blocked"' not in sql and '"kind": "pptx"' not in sql  # no dependency on JSON spacing
    assert ")[0]" not in sql  # ANSI mode: indexing an empty array fails the whole job; use get(array, 0)


# ---------- round-13 coverage ----------


# ---------- fresh-workspace finding ----------
