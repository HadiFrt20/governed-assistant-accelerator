"""Office renderers: deck spec -> .pptx in the corporate template, doc spec -> .docx (blank or filled 8D template)."""
import copy
import io

from docx import Document
from docx.shared import Pt, RGBColor as DRGB
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.util import Inches, Pt as PPt

import os

import config

# Brand palette (config.COLORS): primary = covers and table headers, primary_2 = titles, accent and secondary = chart series
NAVY, ACCENT, GREY, RED = RGBColor(*config.hex_rgb("primary_2")), RGBColor(*config.hex_rgb("secondary")), RGBColor(0x5F, 0x6B, 0x7A), RGBColor(0xD9, 0x3B, 0x2B)
GREEN, INK = RGBColor(*config.hex_rgb("accent")), RGBColor(*config.hex_rgb("primary"))
FONT = config.FONTS.get("office", "Calibri")
WORDMARK = config.TEMPLATE_FILES.get("wordmark", "templates/wordmark.png")   # dark logo, re-inserted on light Office pages
WORDMARK_ON_DARK = os.path.join("static", config.LOGOS["wordmark_on_dark"].lstrip("/"))  # page headers (navy band)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DECK_TEMPLATE = config.TEMPLATE_FILES["pptx"]
D8_TEMPLATE = config.TEMPLATE_FILES["docx_8d"]

CHART_TYPES = {"line": XL_CHART_TYPE.LINE_MARKERS, "bar": XL_CHART_TYPE.COLUMN_CLUSTERED,
               "barh": XL_CHART_TYPE.BAR_CLUSTERED, "pie": XL_CHART_TYPE.PIE}


def _style_run(run, size, color, bold=False):
    run.font.size, run.font.color.rgb, run.font.bold, run.font.name = PPt(size), color, bold, FONT


def _set_text(paragraph, text, size, color, bold=False):
    """Write text into a paragraph and style it; works for empty strings (python-pptx creates no run for "")."""
    for r in list(paragraph.runs)[1:]:
        r._r.getparent().remove(r._r)
    run = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
    run.text = str(text if text is not None else "")
    _style_run(run, size, color, bold)


AI_NOTICE = f"Contenu généré par IA ({config.BRAND}) à partir des données et documents de l'entreprise : à relire avant diffusion."


def _mark_ai(core):
    """Machine-readable AI-generation marking in Office document properties (EU AI Act art. 50 transparency)."""
    core.author = config.BRAND
    core.comments = AI_NOTICE
    core.keywords = "contenu-genere-par-ia"


def _content_slide(prs, template_slide):
    """Clone the template's content slide (accent bar + title + body boxes)."""
    s = prs.slides.add_slide(prs.slide_layouts[6])
    for shp in template_slide.shapes:
        if shp.shape_type == 13:  # picture: re-insert so the image relationship exists on this slide
            s.shapes.add_picture(WORDMARK, shp.left, shp.top, shp.width, shp.height)
        else:
            s.shapes._spTree.insert_element_before(copy.deepcopy(shp._element), "p:extLst")
    return s


def _boxes(slide):
    tbs = [x for x in slide.shapes if x.has_text_frame and x.text_frame.text.strip()]
    body = next(x for x in tbs if x.text_frame.text.strip() == "[contenu]")
    footer = next((x for x in tbs if x.text_frame.text.startswith(f"{config.COMPANY} ·")), None)
    title = next(x for x in tbs if x not in (body, footer))
    return title, body


def _add_table(slide, table, x, y, w):
    cols = [str(c) for c in table.get("columns") or []] or [""]
    rows = [[("" if v is None else v) for v in (list(r) + [""] * len(cols))[:len(cols)]] for r in (table.get("rows") or [])]  # pad / cut ragged rows
    widths = [max([len(c)] + [len(str(r[i])) for r in rows]) for i, c in enumerate(cols)]
    widths = [max(min(v, 60), 5) for v in widths]
    total = sum(widths) or 1
    shape = slide.shapes.add_table(len(rows) + 1, len(cols), x, y, w, Inches(0.42) * (len(rows) + 1))
    tbl = shape.table
    for i, v in enumerate(widths):
        tbl.columns[i].width = int(w * v / total)
    for k, h in enumerate(cols):
        c = tbl.cell(0, k)
        c.fill.solid(); c.fill.fore_color.rgb = INK
        _set_text(c.text_frame.paragraphs[0], h, 12, WHITE, True)
    for r, row in enumerate(rows, start=1):
        for k, v in enumerate(row):
            _set_text(tbl.cell(r, k).text_frame.paragraphs[0], v, 11, NAVY)


def _add_chart(slide, chart, x, y, w, h):
    cd = CategoryChartData()
    cats = [str(c) for c in chart.get("categories") or []]
    cd.categories = cats
    for name, vals in (chart.get("series") or {}).items():
        vals = (list(vals) + [None] * len(cats))[:len(cats)]
        cd.add_series(str(name), [float(v) if isinstance(v, (int, float)) else None for v in vals])
    kind = CHART_TYPES.get(chart.get("type", "line"), XL_CHART_TYPE.LINE_MARKERS)
    ch = slide.shapes.add_chart(kind, x, y, w, h, cd).chart
    ch.has_legend = len(chart["series"]) > 1 or kind == XL_CHART_TYPE.PIE
    if ch.has_legend:
        ch.legend.position, ch.legend.include_in_layout = XL_LEGEND_POSITION.BOTTOM, False
    if kind != XL_CHART_TYPE.PIE and chart.get("y_title"):
        ch.value_axis.has_title = True
        ch.value_axis.axis_title.text_frame.text = chart["y_title"]
    palette = [NAVY, GREEN, ACCENT, GREY]
    for ser, col in zip(ch.plots[0].series, palette):
        if kind == XL_CHART_TYPE.LINE_MARKERS:
            ser.format.line.color.rgb, ser.format.line.width = col, PPt(2.5)
        elif kind != XL_CHART_TYPE.PIE:
            ser.format.fill.solid(); ser.format.fill.fore_color.rgb = col


def build_pptx(spec: dict) -> bytes:
    """spec = {title, subtitle, slides:[{title, bullets[], chart?{type,categories,series{name:[..]},y_title}, table?{columns,rows}, source?}]}"""
    prs = Presentation(DECK_TEMPLATE)
    cover, content_tpl = prs.slides[0], prs.slides[1]
    tb = [s for s in cover.shapes if s.has_text_frame][-1].text_frame
    for para, value in ((tb.paragraphs[0], spec.get("title", "")), (tb.paragraphs[1], spec.get("subtitle", ""))):
        (para.runs[0] if para.runs else para.add_run()).text = str(value or "")
    _mark_ai(prs.core_properties)
    for s in spec.get("slides", []):
        slide = _content_slide(prs, content_tpl)
        title_box, body_box = _boxes(slide)
        title_box.text_frame.word_wrap = True
        _set_text(title_box.text_frame.paragraphs[0], s.get("title", ""), 26 if len(s.get("title", "")) < 52 else 22, NAVY, True)
        visual = "chart" in s or "table" in s
        left, top = Inches(0.6), Inches(1.45)
        body_box.left, body_box.top = left, top
        body_box.width, body_box.height = (Inches(4.5) if visual else Inches(12.1)), Inches(5.4)
        tf = body_box.text_frame
        tf.word_wrap = True
        bullets = [b for b in (s.get("bullets") or []) if str(b).strip()]
        if not bullets:  # no text: remove the template placeholder « [contenu] » instead of leaving it on the slide
            _set_text(tf.paragraphs[0], "", 15, NAVY)
        for j, b in enumerate(bullets):
            p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            p.space_after = PPt(10)
            _set_text(p, b, 15 if visual else 18, NAVY if j == 0 else GREY, j == 0)
        vx, vw = Inches(5.4), Inches(7.4)
        if "chart" in s:
            _add_chart(slide, s["chart"], vx, top, vw, Inches(5.2))
        elif "table" in s:
            _add_table(slide, s["table"], vx, top, vw)
        if s.get("source"):
            src = slide.shapes.add_textbox(Inches(8.4), Inches(7.05), Inches(4.4), Inches(0.3)).text_frame
            _set_text(src.paragraphs[0], "Source : " + str(s["source"]).removeprefix("Source :").removeprefix("Source:").strip(), 9, GREY)
    # drop the 4 placeholder slides that came with the template
    sldIdLst = prs.slides._sldIdLst
    for sid in list(sldIdLst)[1:5]:
        prs.part.drop_rel(sid.rId)
        sldIdLst.remove(sid)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _add_docx_table(doc, table):
    cols = [str(c) for c in table.get("columns") or []] or [""]
    t = doc.add_table(rows=1, cols=len(cols))
    t.style = "Table Grid"
    for i, h in enumerate(cols):
        t.rows[0].cells[i].paragraphs[0].add_run(h).bold = True
    for row in table.get("rows") or []:
        cells = t.add_row().cells
        for i, v in enumerate((list(row) + [""] * len(cols))[:len(cols)]):  # pad / cut ragged rows
            cells[i].text = "" if v is None else str(v)
    return t


def _d_code(heading):
    """'D1 · Équipe', 'D1: Équipe', 'd1. équipe' -> 'D1' ; anything else -> None."""
    m = _re.match(r"\s*(D[1-8])\b", str(heading or ""), _re.I)
    return m.group(1).upper() if m else None


def _ai_footer(doc):
    """Visible AI notice in the footer of every section (AI Act art. 50), whatever the template already carries there."""
    for sec in doc.sections:
        if AI_NOTICE in "\n".join(p.text for p in sec.footer.paragraphs):
            continue
        par = sec.footer.paragraphs[0] if not sec.footer.paragraphs[0].text.strip() else sec.footer.add_paragraph()
        fr = par.add_run(AI_NOTICE)
        fr.font.size, fr.font.color.rgb = Pt(7.5), DRGB(0x5F, 0x6B, 0x7A)


def build_docx(spec: dict) -> bytes:
    """spec = {template: "8D"|null, title, meta?{key: value}, sections:[{heading, paragraphs[], bullets[], table?{columns,rows}}]}
    With template "8D", sections whose heading starts with D1..D8 replace the grey placeholder under that heading."""
    if spec.get("template") == "8D":
        doc = Document(D8_TEMPLATE)
        meta = spec.get("meta", {})
        for row in doc.tables[0].rows:
            key = row.cells[0].text.strip()
            if key in meta:
                row.cells[1].text = str(meta[key])
        by_code = {_d_code(s.get("heading")): s for s in spec.get("sections", []) if _d_code(s.get("heading"))}
        paras = doc.paragraphs
        for i, p in enumerate(paras):
            code = _d_code(p.text) if p.style.name.startswith("Heading") else None
            # only fill the grey placeholder that follows the heading, never the next heading
            if code in by_code and i + 1 < len(paras) and not paras[i + 1].style.name.startswith("Heading"):
                sec, placeholder = by_code[code], paras[i + 1]
                anchor = placeholder
                lines = sec.get("paragraphs", []) + ["• " + b for b in sec.get("bullets", [])]
                placeholder.text = lines[0] if lines else ""
                for r in placeholder.runs:
                    r.font.size, r.font.color.rgb = Pt(10.5), DRGB(0x1F, 0x29, 0x37)
                for line in lines[1:]:
                    np_ = copy.deepcopy(placeholder._p)
                    anchor._p.addnext(np_)
                    from docx.text.paragraph import Paragraph
                    anchor = Paragraph(np_, placeholder._parent)
                    anchor.text = line
                    for r in anchor.runs:
                        r.font.size, r.font.color.rgb = Pt(10.5), DRGB(0x1F, 0x29, 0x37)
                if sec.get("table"):
                    t = _add_docx_table(doc, sec["table"])
                    anchor._p.addnext(t._tbl)
    else:
        doc = Document()
        st = doc.styles["Normal"]; st.font.size = Pt(10.5)
        st.font.name = FONT
        for sty, size in (("Title", 22), ("Heading 1", 13)):
            doc.styles[sty].font.name, doc.styles[sty].font.size, doc.styles[sty].font.color.rgb = FONT, Pt(size), DRGB(*config.hex_rgb("primary_2"))
        hdr = doc.sections[0].header.paragraphs[0]
        from docx.shared import Cm
        hdr.add_run().add_picture(WORDMARK, height=Cm(0.55))
        hr = hdr.add_run(f"\t{spec.get('doc_type', 'Document')} · {config.SITE} · Diffusion interne")
        hr.font.size, hr.font.color.rgb = Pt(8), DRGB(0x5F, 0x6B, 0x7A)
        doc.add_heading(spec.get("title", ""), 0)
        import datetime
        meta = doc.add_paragraph()
        meta.add_run(f"{spec.get('doc_type', 'Document')} · {datetime.date.today():%d/%m/%Y} · Rédacteur : {spec.get('author', '')}").font.color.rgb = DRGB(0x5F, 0x6B, 0x7A)
        for sec in spec.get("sections", []):
            doc.add_heading(sec.get("heading", ""), 1)
            for para in sec.get("paragraphs", []):
                doc.add_paragraph(para)
            for b in sec.get("bullets", []):
                doc.add_paragraph(b, style="List Bullet")
            if sec.get("table"):
                _add_docx_table(doc, sec["table"])
    _mark_ai(doc.core_properties)
    _ai_footer(doc)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------- interactive HTML pages (brand design system) ----------
import base64 as _b64
import os
import datetime as _dt
import html as _html
import json as _json
import re as _re

_DS = os.path.join(os.path.dirname(__file__), "design")


def _ds(name):
    return open(os.path.join(_DS, name), encoding="utf-8").read()


def _sanitize(body):
    """body_html is layout only (sections and empty containers): no CSS, no script, no handlers, no embedded or external content.
    The page's behaviour lives in `script` (sandboxed, no network: see GENERATED_PAGE_CSP in app.py)."""
    body = _re.sub(r"(?is)<(style|script|iframe|object|embed|noscript|template|svg|math)\b.*?</\1\s*>", "", body)
    body = _re.sub(r"(?is)<(link|meta|base|script|iframe|object|embed|form|input|button|frame|frameset)\b[^>]*>", "", body)
    body = _re.sub(r"(?is)</?(html|head|body|title|form|frameset)\b[^>]*>", "", body)
    body = _re.sub(r"(?is)[\s/]on[a-z]+\s*=\s*(\"[^\"]*\"|'[^']*'|[^\s>]+)", " ", body)              # event handlers
    body = _re.sub(r"(?is)[\s/](href|src|srcset|action|formaction|xlink:href|data|poster)\s*=\s*(\"[^\"]*\"|'[^']*'|[^\s>]+)", " ", body)  # no links or sources

    def keep_width(m):
        decl = [d.strip() for d in m.group(2).split(";") if _re.match(r"\s*(width|--c|--cb)\s*:", d)]
        return f' style="{"; ".join(decl)}"' if decl else ""
    return _re.sub(r'(?is)\sstyle=(["\'])(.*?)\1', keep_width, body)


def build_html(spec):
    title, sub = spec.get("title", config.COMPANY), spec.get("subtitle", "")
    tabs = spec.get("tabs") or []
    logo = _b64.b64encode(open(os.path.join(os.path.dirname(__file__), WORDMARK_ON_DARK), "rb").read()).decode()
    data = spec.get("data")
    script = spec.get("script", "")
    script = _re.sub(r"(?is)</script", "<\\/script", script)
    tab_html = "".join(f'<button class="em-tab" data-tab="{_html.escape(t["id"])}">{_html.escape(t["label"])}</button>' for t in tabs)
    today = _dt.date.today().strftime("%d/%m/%Y")
    eyebrow = spec.get("eyebrow") or f"{config.COMPANY} · {config.SITE}"
    # JSON inside <script>: escape <, > and & so a value can never close the tag or open a comment
    data_js = _json.dumps(data, ensure_ascii=False).replace("<", "\\u003c").replace(">", "\\u003e").replace("&", "\\u0026") if data is not None else "null"
    return f"""<!doctype html>
<html lang="fr" data-ai-generated="true"><head><meta charset="utf-8"><meta name="ai-generated" content="{_html.escape(AI_NOTICE)}"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>{_html.escape(title)} · {_html.escape(config.COMPANY)}</title>
<link rel="preconnect" href="https://fonts.googleapis.com"><link href="https://fonts.googleapis.com/css2?family={config.FONTS.get("heading", "Titillium Web").replace(" ", "+")}:wght@400;600;700&display=swap" rel="stylesheet">
<style>{config.rebrand(_ds("ds.css"))}</style></head>
<body>
<header class="em-top"><div class="in">
<div class="bar"><img src="data:image/png;base64,{logo}" alt="{_html.escape(config.COMPANY)}"><span class="org">{_html.escape(spec.get("org") or config.ORG)}</span><button class="em-theme" type="button">Clair / sombre</button></div>
<div class="hero"><span class="em-eyebrow">{_html.escape(eyebrow)}</span><h1>{_html.escape(title)}</h1>{f'<p class="sub">{_html.escape(sub)}</p>' if sub else ''}<div class="meta">Publié le {today} par {_html.escape(config.BRAND)} · document interne</div></div>
{f'<nav class="em-tabs">{tab_html}</nav>' if tabs else ''}</div></header>
<main>{_sanitize(spec.get("body_html", ""))}</main>
<footer class="em-foot"><i></i>Document interne {_html.escape(config.COMPANY)} · généré par {_html.escape(config.BRAND)} à partir de données gouvernées par Unity Catalog · à vérifier avant diffusion</footer>
<div class="em-scrim" id="em-scrim"></div>
<aside class="em-drawer" id="em-drawer" aria-live="polite"><div class="h"><div style="flex:1"><span class="em-eyebrow dark ey"></span><h3></h3></div><button class="x" onclick="EM.close()" aria-label="Fermer">×</button></div><div class="c"></div></aside>
<script>{_ds("ds.js")}</script>
<script>const DATA = {data_js};</script>
<script>document.addEventListener("DOMContentLoaded", () => {{ try {{
{script}
}} catch (e) {{ console.error(e); document.querySelector("main").insertAdjacentHTML("afterbegin", '<div class="em-callout no"><div>Erreur d\\'affichage : ' + EM.esc(e.message) + '</div></div>'); }} }});</script>
</body></html>"""
