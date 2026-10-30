"""Read Office and PDF documents directly in the app (no LLM call): text and tables, ready to be put in the model context."""
import io

from docx import Document
from pptx import Presentation

MAX_CHARS = 120_000


def _docx(data):
    d = Document(io.BytesIO(data))
    out = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        out.append("\n".join(" | ".join(c.text.strip() for c in r.cells) for r in t.rows))
    return "\n".join(out)


def _pptx(data):
    out = []
    for i, s in enumerate(Presentation(io.BytesIO(data)).slides, 1):
        parts = []
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                parts.append(sh.text_frame.text.strip())
            if getattr(sh, "has_table", False) and sh.has_table:
                parts.append("\n".join(" | ".join(c.text for c in r.cells) for r in sh.table.rows))
        if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip():
            parts.append("Notes : " + s.notes_slide.notes_text_frame.text.strip())
        out.append(f"[Diapositive {i}]\n" + "\n".join(parts))
    return "\n\n".join(out)


MAX_PAGES = 300


def _xlsx(data, max_rows=2000):
    """Spreadsheet as text, read straight from the OOXML parts (no extra dependency): one line per row, ' | ' between cells."""
    import re
    import zipfile
    from xml.etree import ElementTree as ET
    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        names = z.namelist()
        shared = []
        if "xl/sharedStrings.xml" in names:
            for si in ET.fromstring(z.read("xl/sharedStrings.xml")).findall("m:si", ns):
                shared.append("".join(t.text or "" for t in si.iter(f"{{{ns['m']}}}t")))
        sheets = sorted((n for n in names if re.match(r"xl/worksheets/sheet\d+\.xml$", n)), key=lambda n: int(re.findall(r"\d+", n)[-1]))
        out = []
        for i, name in enumerate(sheets, 1):
            out.append(f"[Feuille {i}]")
            for row in ET.fromstring(z.read(name)).iter(f"{{{ns['m']}}}row"):
                cells = []
                for c in row.findall("m:c", ns):
                    v = c.find("m:v", ns)
                    if c.get("t") == "s" and v is not None and v.text and v.text.isdigit() and int(v.text) < len(shared):
                        cells.append(shared[int(v.text)])
                    elif c.get("t") == "inlineStr":
                        cells.append("".join(t.text or "" for t in c.iter(f"{{{ns['m']}}}t")))
                    else:
                        cells.append(v.text if v is not None and v.text else "")
                out.append(" | ".join(cells))
                if len(out) > max_rows:
                    out.append("[… lignes suivantes non lues]")
                    return "\n".join(out)
    return "\n".join(out)


def _pdf(data):
    from pypdf import PdfReader
    pages = PdfReader(io.BytesIO(data)).pages
    return "\n\n".join(f"[Page {i}]\n{p.extract_text() or ''}" for i, p in enumerate(pages[:MAX_PAGES], 1))


MAX_UNZIPPED = 200 * 1024 * 1024


def _check_zip(data):
    """Office files are zip archives: refuse decompression bombs before parsing."""
    import zipfile
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        total = sum(i.file_size for i in z.infolist())
    if total > MAX_UNZIPPED or total > 100 * max(len(data), 1):
        raise ValueError("archive trop volumineuse une fois décompressée")


def extract(name, data):
    ext = name.rsplit(".", 1)[-1].lower()
    if ext in ("docx", "pptx", "xlsx"):
        _check_zip(data)
    if ext == "docx":
        text = _docx(data)
    elif ext == "pptx":
        text = _pptx(data)
    elif ext == "xlsx":
        text = _xlsx(data)
    elif ext == "pdf":
        text = _pdf(data)
    elif ext in ("txt", "csv", "md"):
        text = data.decode("utf-8", errors="replace")
    else:
        return None
    text = text.replace("\x00", "")  # Postgres TEXT/JSONB reject NUL characters
    return text[:MAX_CHARS] + ("\n[… document tronqué]" if len(text) > MAX_CHARS else "")
